"""Build the NetGuard AI project report (VIT BCSE308P template) as a Word document.

Usage (from the project root, after ml.train and ml.analyze_dataset have run):
    python -m report.build_report

The script reads reports/metrics.json and reports/dataset_analysis.json, draws the
diagrams, and writes docs/report/NetGuard_AI_Project_Report.docx. Open the file in
Word and update fields (Ctrl+A, F9) to fill the Contents and the lists of figures and
tables; report/update_word.ps1 does this automatically and also exports a PDF.
"""

import json
from pathlib import Path

from docx.shared import Inches, Pt

from report import content_part1 as p1
from report import content_part2 as p2
from report import content_part3 as p3
from report.diagrams import draw_all
from report.docx_helpers import Report

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "report" / "assets"
DIAGRAMS = ROOT / "report" / "assets" / "diagrams"
FIGURES = ROOT / "reports" / "figures"
OUTPUT = ROOT / "docs" / "report" / "NetGuard_AI_Project_Report.docx"

TITLE = "NetGuard AI: Malware Detection Using AI-Based Network Traffic Analysis"
COURSE = "BCSE308P – Computer Networks"
DEGREE = ("Bachelor of Computer Science and Engineering (Artificial Intelligence and Data Engineering) / "
          "Bachelor of Computer Science and Engineering (Cyber Security)")


def cover(R: Report) -> None:
    R.spacer(10)
    R.centered(TITLE, size=18, bold=True, space_after=8)
    R.centered("BCSE308P – COURSE PROJECT", size=14, bold=True, space_after=22)
    R.centered("Submitted in partial fulfillment of the requirements for the degree of", size=12, italic=True,
               space_after=22)
    R.centered("Bachelor of Computer Science and Engineering (Artificial Intelligence and Data Engineering) / "
               "Bachelor of Computer Science and Engineering (Cyber Security)", size=16, bold=True,
               space_after=14, highlight=True)
    R.centered("in", size=12, bold=True, space_after=14)
    R.centered("Department of Cybersecurity and Intelligent Systems", size=16, bold=True, space_after=30)
    R.centered("by", size=12, italic=True, space_after=8)
    R.centered("<Student Names>", size=14, bold=True, space_after=4, highlight=True)
    R.centered("<Register Numbers>", size=14, bold=True, space_after=24, highlight=True)
    p = R.doc.add_paragraph()
    p.alignment = 1
    p.add_run().add_picture(str(ASSETS / "vit_logo.png"), width=Inches(2.9))
    R.spacer(10)
    R.centered("School of Computer Science Engineering and Information Systems [SCORE]", size=12, bold=True,
               space_after=4)
    R.centered("VIT, Vellore", size=12, bold=True)
    R.page_break()


def declaration(R: Report) -> None:
    R.centered("November, 2026", size=12, space_after=60, highlight=True)
    R.front_heading("DECLARATION")
    R.para(
        f"I hereby declare that the {COURSE} Course Project entitled \"{TITLE}\" submitted by me, for the award of "
        f"the degree of {DEGREE}, School of Computer Science Engineering and Information Systems to VIT is a "
        "record of bonafide work carried out by me under the supervision of <Faculty Guide Name, Designation, "
        "SCORE, VIT, Vellore>.", indent=False)
    R.para(
        "I further declare that the work reported in this project has not been submitted and will not be "
        "submitted, either in part or in full, for the award of any other degree or diploma in this institute "
        "or any other institute or university.", indent=True)
    R.spacer(40)
    R.para("Place: Vellore", indent=False)
    R.para("Date:", indent=False)
    R.spacer(60)
    p = R.doc.add_paragraph()
    p.alignment = 2
    run = p.add_run("Signature of the Candidate")
    run.bold = True
    R.page_break()


def certificate(R: Report) -> None:
    R.front_heading("CERTIFICATE")
    R.para(
        f"This is to certify that the {COURSE} Course Project entitled \"{TITLE}\" submitted by <Student Name & "
        f"Reg. No.>, SCORE, VIT, for the award of the degree of {DEGREE}, School of Computer Science Engineering "
        "and Information Systems, is a record of bonafide work carried out under my supervision during the "
        "period <start date> to <end date>, as per the VIT code of academic and research ethics.", indent=True)
    R.para(
        "The contents of this report have not been submitted and will not be submitted either in part or in "
        "full, for the award of any other degree or diploma in this institute or any other institute or "
        "university. The project fulfills the requirements and regulations of the University and in my "
        "opinion meets the necessary standards for submission.", indent=True)
    R.spacer(40)
    R.para("Place: Vellore", indent=False)
    R.para("Date:", indent=False)
    R.spacer(60)
    p = R.doc.add_paragraph()
    p.alignment = 2
    run = p.add_run("Signature of the VIT-SCORE - Guide")
    run.bold = True
    R.page_break()


def acknowledgement(R: Report) -> None:
    R.front_heading("ACKNOWLEDGEMENT", "Acknowledgement")
    R.para(
        "I express my sincere gratitude to my guide, <Faculty Guide Name>, School of Computer Science Engineering "
        "and Information Systems, VIT, Vellore, for the guidance, encouragement and constructive feedback given "
        "throughout this project.", indent=False)
    R.para(
        "I thank the Head of the Department of Cybersecurity and Intelligent Systems and the Dean of SCORE for "
        "providing the academic environment and facilities that made this work possible.", indent=True)
    R.para(
        "I am grateful to the Canadian Institute for Cybersecurity, University of New Brunswick, for publishing "
        "the CIC-IDS2017 dataset, and to the developers of the open-source tools used in this project, "
        "including pandas, scikit-learn, FastAPI, React and Vite.", indent=True)
    R.para("Finally, I thank my family and friends for their constant support.", indent=True)
    R.spacer(30)
    p = R.doc.add_paragraph()
    p.alignment = 2
    p.add_run("<Student Name>").bold = True
    R.page_break()


def executive_summary(R: Report, A: dict, M: dict) -> str:
    rf = next(m for m in M["models"] if m["name"] == "Random Forest")
    text = (
        "NetGuard AI is a web application that detects malicious network traffic by applying machine learning "
        f"to network flow records. It is built on the public CIC-IDS2017 dataset, whose {A['raw_rows']:,} flows "
        "were analysed in detail. The analysis found "
        f"{A['raw_rows'] - A['dedup_rows']:,} exact duplicate rows, infinite and missing rate values caused by "
        "zero-duration flows, protocol sentinel values, eight constant features and an extremely uneven "
        "distribution of attack types. A preprocessing pipeline removes duplicates before a stratified "
        "500,000-flow sample is split into training, validation and test sets. Random Forest, Decision Tree "
        "and Logistic Regression were trained and compared; the Random Forest, selected on the validation set, "
        f"achieved {100 * rf['accuracy']:.2f}% accuracy, {100 * rf['precision']:.2f}% precision, "
        f"{100 * rf['recall']:.2f}% recall and an F1-score of {100 * rf['f1_score']:.2f}% on 100,000 unseen "
        "flows. A per-attack evaluation showed near-perfect detection of denial-of-service, port-scan and "
        "brute-force attacks but weak detection of the rarest labels, a direct consequence of the dataset. The "
        "model is served through a validated FastAPI backend and a React dashboard that accepts CSV uploads, "
        "shows predictions and genuine evaluation metrics, and exports the results."
    )
    R.spacer(120)
    R.front_heading("Executive Summary", "Executive Summary")
    R.para(text, indent=False)
    R.page_break()
    return text


def abbreviations(R: Report) -> None:
    R.front_heading("List of Abbreviations", "Abbreviations")
    items = [
        ("ACK", "Acknowledgement (TCP flag)"), ("API", "Application Programming Interface"),
        ("AUC", "Area Under the Curve"), ("AP", "Average Precision"),
        ("CIC", "Canadian Institute for Cybersecurity"), ("CIC-IDS2017", "CIC Intrusion Detection Evaluation Dataset 2017"),
        ("CORS", "Cross-Origin Resource Sharing"), ("CSV", "Comma-Separated Values"),
        ("CWR / CWE", "Congestion Window Reduced (TCP flag; CWE in the dataset)"), ("DDoS", "Distributed Denial of Service"),
        ("DNS", "Domain Name System"), ("DoS", "Denial of Service"), ("DT", "Decision Tree"),
        ("ECE", "ECN-Echo (TCP flag)"), ("ECN", "Explicit Congestion Notification"), ("FIN", "Finish (TCP flag)"),
        ("FN", "False Negative"), ("FP", "False Positive"), ("FPR", "False Positive Rate"),
        ("FTP", "File Transfer Protocol"), ("HIDS", "Host-based Intrusion Detection System"),
        ("HTTP(S)", "Hypertext Transfer Protocol (Secure)"), ("IAT", "Inter-Arrival Time"),
        ("IDS", "Intrusion Detection System"), ("IETF", "Internet Engineering Task Force"),
        ("IP", "Internet Protocol"), ("IPFIX", "IP Flow Information Export"), ("IR", "Imbalance Ratio"),
        ("JSON", "JavaScript Object Notation"), ("LR", "Logistic Regression"), ("MDI", "Mean Decrease in Impurity"),
        ("ML", "Machine Learning"), ("NaN", "Not a Number (missing value)"),
        ("NIDS", "Network-based Intrusion Detection System"), ("NTP", "Network Time Protocol"),
        ("OAS", "OpenAPI Specification"), ("PCAP", "Packet Capture (file format)"), ("PSH", "Push (TCP flag)"),
        ("REST", "Representational State Transfer"), ("RF", "Random Forest"), ("RFC", "Request for Comments"),
        ("ROC", "Receiver Operating Characteristic"), ("RST", "Reset (TCP flag)"),
        ("SMOTE", "Synthetic Minority Over-sampling Technique"), ("SQL", "Structured Query Language"),
        ("SSH", "Secure Shell"), ("SYN", "Synchronise (TCP flag)"), ("TCP", "Transmission Control Protocol"),
        ("TLS", "Transport Layer Security"), ("TN", "True Negative"), ("TP", "True Positive"),
        ("UDP", "User Datagram Protocol"), ("UI", "User Interface"), ("UNB", "University of New Brunswick"),
        ("URG", "Urgent (TCP flag)"), ("WCAG", "Web Content Accessibility Guidelines"), ("XSS", "Cross-Site Scripting"),
    ]
    _two_column_list(R, items, 1.7)
    R.page_break()


def symbols(R: Report) -> None:
    R.front_heading("Symbols and Notations", "Symbols and Notations")
    items = [
        ("x, xᵢ", "Feature vector of a flow / of the i-th flow"), ("xᵢⱼ", "Value of feature j for flow i"),
        ("y, yᵢ", "Class label (0 = benign, 1 = malicious)"), ("ŷ", "Predicted label"),
        ("n, N", "Number of flows"), ("N꜀", "Number of flows of class c"), ("p", "Number of features (69)"),
        ("T", "Number of trees in the Random Forest (100)"), ("m", "Features considered at each split (⌊√p⌋ = 8)"),
        ("hₜ(x)", "Prediction of tree t"), ("pₖ", "Share of class k at a tree node"), ("G", "Gini impurity"),
        ("ΔG", "Impurity decrease of a split"), ("Impⱼ", "Importance of feature j"),
        ("σ(·)", "Sigmoid function"), ("w, b", "Weights and bias of Logistic Regression"),
        ("μⱼ, σⱼ", "Mean and standard deviation of feature j (training set)"), ("zᵢⱼ", "Standardised feature value"),
        ("r", "Pearson correlation coefficient"), ("TP, TN, FP, FN", "Confusion-matrix counts"),
        ("F₁", "F1-score"), ("µs", "Microsecond (unit of Flow Duration)"),
    ]
    _two_column_list(R, items, 1.7)


def _two_column_list(R: Report, items, first_width: float) -> None:
    table = R.doc.add_table(rows=0, cols=2)
    for a, b in items:
        cells = table.add_row().cells
        for cell, text in zip(cells, (a, b)):
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing = 1.15
            run = par.add_run(text)
            run.font.size = Pt(12)
            run.font.name = "Times New Roman"
    for row in table.rows:
        row.cells[0].width = Inches(first_width)
        row.cells[1].width = Inches(6.0 - first_width)


def build() -> Path:
    A = json.loads((ROOT / "reports" / "dataset_analysis.json").read_text())
    M = json.loads((ROOT / "reports" / "metrics.json").read_text())
    draw_all(DIAGRAMS, A, M, p3.SCHEDULE)

    R = Report()
    R.front_matter_numbering()
    cover(R)
    declaration(R)
    certificate(R)
    acknowledgement(R)
    summary = executive_summary(R, A, M)
    R.toc("CONTENTS", "C")
    R.page_break()
    R.toc("List of Figures", "F", ("Figure No.", "Title", "Page No."))
    R.page_break()
    R.toc("List of Tables", "T", ("Table No.", "Title", "Page No."))
    R.page_break()
    abbreviations(R)
    symbols(R)

    R.start_main_section()
    p1.chapter1(R, A, M, DIAGRAMS)
    p1.chapter2(R, A, M, DIAGRAMS)
    p1.chapter3(R, A, M, DIAGRAMS)
    p2.chapter4(R, A, M, DIAGRAMS, FIGURES)
    p3.chapter5(R, A, M, DIAGRAMS)
    p3.chapter6(R, A, M, ASSETS)
    p3.chapter7(R, A, M, FIGURES)
    p3.chapter8(R, A, M)
    p3.chapter9(R)

    R.start_appendix_section()
    p3.appendix(R, A, M)
    R.save(OUTPUT)
    print(f"Saved {OUTPUT}")
    print(f"Executive summary: {len(summary.split())} words")
    return OUTPUT


if __name__ == "__main__":
    build()
