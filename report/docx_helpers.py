"""Helpers that apply the VIT BCSE308P report template rules with python-docx.

Template rules implemented here:
  * Letter page, margins left 1.5", right/top/bottom 1"
  * Times New Roman 12 pt, justified, 1.5 line spacing
  * Chapter heading 14 pt BOLD CAPS, sub-headings 12 pt (no caps)
  * First paragraph of a chapter not indented, later paragraphs indented one tab
  * Check-mark bullets (numbered lists only for types / classifications)
  * "Fig. C.N Caption" below figures, "Table C.N Caption" above tables (10 pt, title case)
  * Equations written with Word's equation tool, numbered (C.N) at the right margin
  * Roman page numbers for the front matter (hidden on the cover), arabic from
    Chapter 1 to the References, no page numbers in the appendices
"""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
MONO = "Courier New"  # plain monospaced font for code listings
TEXT_WIDTH_IN = 6.0  # 8.5" page - 1.5" left - 1" right
M_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


# ---------------------------------------------------------------- equations (OMML)
def m_run(text: str, plain: bool = False) -> str:
    style = '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain else ""
    return f'<m:r>{style}<m:t xml:space="preserve">{escape(text)}</m:t></m:r>'


def m_text(text: str) -> str:
    """Upright words inside an equation (e.g. 'Precision')."""
    return m_run(text, plain=True)


def m_frac(num: str, den: str) -> str:
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


def m_sub(base: str, sub: str) -> str:
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>"


def m_sup(base: str, sup: str) -> str:
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{sup}</m:sup></m:sSup>"


def m_subsup(base: str, sub: str, sup: str) -> str:
    return f"<m:sSubSup><m:e>{base}</m:e><m:sub>{sub}</m:sub><m:sup>{sup}</m:sup></m:sSubSup>"


def m_sum(sub: str, sup: str | None, body: str, char: str = "∑") -> str:
    """n-ary sum; pass sup=None for a sum with only a lower limit."""
    hide = '<m:supHide m:val="1"/>' if sup is None else ""
    return (f'<m:nary><m:naryPr><m:chr m:val="{char}"/><m:limLoc m:val="undOvr"/>{hide}</m:naryPr>'
            f"<m:sub>{sub}</m:sub><m:sup>{sup or ''}</m:sup><m:e>{body}</m:e></m:nary>")


def m_paren(body: str, left: str = "(", right: str = ")") -> str:
    return (f'<m:d><m:dPr><m:begChr m:val="{escape(left)}"/><m:endChr m:val="{escape(right)}"/>'
            f"</m:dPr><m:e>{body}</m:e></m:d>")


def m_sqrt(body: str) -> str:
    return f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{body}</m:e></m:rad>'


def m_acc_hat(text: str) -> str:
    """Upright text with a hat accent, e.g. an estimated quantity."""
    return f'<m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e>{m_text(text)}</m:e></m:acc>'


def m_bar(body: str) -> str:
    return f'<m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e>{body}</m:e></m:acc>'


# ---------------------------------------------------------------- low-level XML
def _set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _add_field(paragraph, instr: str, placeholder: str = "") -> None:
    """Insert a complex field (TOC, PAGE, ...) that Word fills in on update."""
    def fld_char(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    run = paragraph.add_run()
    run._r.append(fld_char("begin"))
    run = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run._r.append(instr_el)
    run = paragraph.add_run()
    run._r.append(fld_char("separate"))
    if placeholder:
        paragraph.add_run(placeholder)
    run = paragraph.add_run()
    run._r.append(fld_char("end"))


def _font(run, size: float = 12, bold: bool = False, italic: bool = False,
          name: str = FONT, color: RGBColor | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), name)


# ---------------------------------------------------------------- report builder
class Report:
    def __init__(self) -> None:
        self.doc = Document()
        self.chapter_no = 0
        self.fig_no = 0
        self.table_no = 0
        self.eq_no = 0
        self.first_paragraph = True
        self.entries = {"C": [], "F": [], "T": []}   # contents, figures, tables
        self.refs: dict[str, str] = {}               # named cross-references, e.g. "Table 4.6"
        self._placeholders = {}
        self._bookmark_id = 0
        self._setup_styles()
        self._setup_page(self.doc.sections[0])
        self._bullet_num_id = self._create_bullet_numbering()

    # ------------------------------------------------------------ list entries
    def _mark(self, paragraph, kind: str, level: int, number: str, title: str) -> None:
        """Bookmark a heading/caption and register it for the Contents or a list."""
        self._bookmark_id += 1
        name = f"_ng{self._bookmark_id:04d}"
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._bookmark_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._bookmark_id))
        p_pr = paragraph._p.pPr
        if p_pr is not None:
            p_pr.addnext(start)
        else:
            paragraph._p.insert(0, start)
        paragraph._p.append(end)
        self.entries[kind].append((level, number, title, name))

    def finalize_lists(self) -> None:
        """Write the Contents / List of Figures / List of Tables entries.

        Each entry ends with a PAGEREF field to the bookmark of its heading or
        caption; Word fills in the page numbers when fields are updated."""
        for kind, placeholder in self._placeholders.items():
            for level, number, title, bookmark in self.entries[kind]:
                p = placeholder.insert_paragraph_before()
                p.style = self.doc.styles[f"toc {level if kind == 'C' else 1}"]
                if kind != "C":
                    p.paragraph_format.left_indent = Inches(0.8)
                    p.paragraph_format.first_line_indent = Inches(-0.8)
                    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.8), WD_TAB_ALIGNMENT.LEFT)
                bold = kind == "C" and level == 1
                _font(p.add_run(f"{number}\t{title}\t"), 12, bold=bold)
                _add_field(p, f"PAGEREF {bookmark} \\h", "0")
                for r in p.runs:
                    if r.font.name != FONT:
                        _font(r, 12, bold=bold)
            placeholder._p.getparent().remove(placeholder._p)

    # ------------------------------------------------------------ setup
    def _setup_page(self, section) -> None:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.footer_distance = Inches(0.5)

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        fmt = normal.paragraph_format
        fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        fmt.space_after = Pt(6)
        fmt.space_before = Pt(0)
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # Word's built-in "toc N" styles format the Contents and the lists of figures/tables:
        # number, tab, title, dotted right tab, page number.
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_TAB_LEADER
        for level, (indent, title_tab) in enumerate([(0.0, 0.55), (0.3, 1.0), (0.6, 1.5)], start=1):
            name = f"toc {level}"
            try:
                st = styles[name]
            except KeyError:
                st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                st.element.set(qn("w:styleId"), f"TOC{level}")
            st.base_style = normal
            st.font.name = FONT
            st.font.size = Pt(12)
            pf = st.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.left_indent = Inches(indent + 0.55 if level == 1 else title_tab)
            pf.first_line_indent = Inches(-(0.55 if level == 1 else title_tab - indent))
            pf.space_after = Pt(3)
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.tab_stops.add_tab_stop(Inches(title_tab), WD_TAB_ALIGNMENT.LEFT)
            pf.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    def _create_bullet_numbering(self) -> int:
        """A check-mark bullet (Wingdings 0xFC) as the template asks."""
        numbering = self.doc.part.numbering_part.element
        abstract_id = 90
        abstract = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
            '<w:multiLevelType w:val="hybridMultilevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
            '<w:lvlText w:val=""/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr>'
            '<w:rPr><w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/></w:rPr>'
            '</w:lvl></w:abstractNum>'
        )
        numbering.insert(0, abstract)
        num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{abstract_id}">'
                        f'<w:abstractNumId w:val="{abstract_id}"/></w:num>')
        numbering.append(num)
        self._next_num_id = abstract_id + 1
        self._numbered_abstract = None
        return abstract_id

    def _new_numbered_list(self) -> int:
        """Decimal numbering (1., 2., ...) that restarts for every list."""
        numbering = self.doc.part.numbering_part.element
        if self._numbered_abstract is None:
            self._numbered_abstract = 80
            numbering.insert(0, parse_xml(
                f'<w:abstractNum {nsdecls("w")} w:abstractNumId="80">'
                '<w:multiLevelType w:val="hybridMultilevel"/>'
                '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
                '<w:lvlText w:val="%1."/><w:lvlJc w:val="left"/>'
                '<w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl></w:abstractNum>'))
        num_id = self._next_num_id
        self._next_num_id += 1
        numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="80"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
        return num_id

    # ------------------------------------------------------------ page numbering
    def _footer_page_number(self, section, fmt: str, start: int | None,
                            hide_first: bool = False) -> None:
        sect_pr = section._sectPr
        pg = sect_pr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sect_pr.append(pg)
        pg.set(qn("w:fmt"), fmt)
        if start is not None:
            pg.set(qn("w:start"), str(start))
        section.footer.is_linked_to_previous = False
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_field(p, "PAGE")
        for r in p.runs:
            _font(r, 12)
        if hide_first:
            section.different_first_page_header_footer = True
            section.first_page_footer.is_linked_to_previous = False

    def front_matter_numbering(self) -> None:
        self._footer_page_number(self.doc.sections[0], "lowerRoman", 1, hide_first=True)

    def start_main_section(self) -> None:
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup_page(section)
        section.different_first_page_header_footer = False
        self._footer_page_number(section, "decimal", 1)

    def start_appendix_section(self) -> None:
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup_page(section)
        section.footer.is_linked_to_previous = False
        for p in section.footer.paragraphs:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
        sect_pr = section._sectPr
        pg = sect_pr.find(qn("w:pgNumType"))
        if pg is not None:
            sect_pr.remove(pg)

    # ------------------------------------------------------------ text blocks
    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def centered(self, text: str, size: float = 12, bold: bool = False, italic: bool = False,
                 space_after: float = 6, highlight: bool = False):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(text)
        _font(r, size, bold, italic)
        return p

    def spacer(self, points: float = 12) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(points)

    def front_heading(self, text: str, toc_entry: str | None = None) -> None:
        p = self.centered(text, size=14, bold=True, space_after=18)
        if toc_entry:
            self._mark(p, "C", 1, "", toc_entry)

    def chapter(self, title: str, new_page: bool = True) -> None:
        """'CHAPTER N' + TITLE, 14 pt bold caps, centred, on a new page."""
        self.chapter_no += 1
        self.fig_no = self.table_no = self.eq_no = 0
        self.first_paragraph = True
        if new_page and self.chapter_no > 1:
            self.page_break()
        p = self.centered(f"CHAPTER {self.chapter_no}", size=14, bold=True, space_after=6)
        p.paragraph_format.keep_with_next = True
        p = self.centered(title.upper(), size=14, bold=True, space_after=18)
        p.paragraph_format.keep_with_next = True
        self._mark(p, "C", 1, str(self.chapter_no), title.upper())

    def unnumbered_chapter(self, title: str) -> None:
        """For the References chapter, which the template numbers as Chapter 9."""
        self.chapter(title)

    def heading(self, number: str, text: str, level: int = 2) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{number} {text}")
        _font(r, 12, bold=True, italic=(level == 3))
        if level <= 3:
            self._mark(p, "C", level, number, text)

    def para(self, text: str, indent: bool | None = None) -> None:
        """Body paragraph. **bold** and *italic* inline markers are supported."""
        p = self.doc.add_paragraph()
        use_indent = (not self.first_paragraph) if indent is None else indent
        if use_indent:
            p.paragraph_format.first_line_indent = Inches(0.5)
        self.first_paragraph = False
        self._add_inline(p, text)

    def _add_inline(self, p, text: str, size: float = 12) -> None:
        import re
        tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|~[^~]+~)", text)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("~"):  # ~k~ -> subscript k (e.g. p~k~ for p with subscript k)
                run = p.add_run(tok[1:-1])
                _font(run, size)
                run.font.subscript = True
            elif tok.startswith("**"):
                _font(p.add_run(tok[2:-2]), size, bold=True)
            elif tok.startswith("`"):
                _font(p.add_run(tok[1:-1]), size - 1.5, name=MONO)
            elif tok.startswith("*"):
                _font(p.add_run(tok[1:-1]), size, italic=True)
            else:
                _font(p.add_run(tok), size)

    def bullets(self, items: list[str], numbered: bool = False) -> None:
        num_id = self._new_numbered_list() if numbered else self._bullet_num_id
        for item in items:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p_pr = p._p.get_or_add_pPr()
            num_pr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/>'
                               f'<w:numId w:val="{num_id}"/></w:numPr>')
            p_pr.insert(0, num_pr)
            self._add_inline(p, item)
        self.first_paragraph = False

    # ------------------------------------------------------------ figures / tables
    def _caption(self, label: str, caption: str, identifier: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        _font(p.add_run(f"{label} "), 10, bold=True)
        _font(p.add_run(caption), 10)
        number = label.split(" ", 1)[1]
        self._mark(p, identifier, 1, number, caption)
        return p

    def figure(self, path: Path | str, caption: str, width_in: float = TEXT_WIDTH_IN) -> None:
        self.fig_no += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.add_run().add_picture(str(path), width=Inches(min(width_in, TEXT_WIDTH_IN)))
        self._caption(f"Fig. {self.chapter_no}.{self.fig_no}", caption, "F")

    @property
    def next_fig(self) -> str:
        return f"Fig. {self.chapter_no}.{self.fig_no + 1}"

    @property
    def next_table(self) -> str:
        return f"Table {self.chapter_no}.{self.table_no + 1}"

    @property
    def next_eq(self) -> str:
        return f"Eqn. {self.chapter_no}.{self.eq_no + 1}"

    def eq_after(self, k: int) -> str:
        """The equation k positions after the next one."""
        return f"Eqn. {self.chapter_no}.{self.eq_no + 1 + k}"

    def fig_after(self, k: int) -> str:
        return f"Fig. {self.chapter_no}.{self.fig_no + 1 + k}"

    @property
    def prev_table(self) -> str:
        return f"Table {self.chapter_no}.{self.table_no}"

    @property
    def prev_fig(self) -> str:
        return f"Fig. {self.chapter_no}.{self.fig_no}"

    def table(self, caption: str, headers: list[str], rows: list[list], widths: list[float] | None = None,
              font_size: float = 10, align: list[str] | None = None) -> None:
        self.table_no += 1
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(8)
        cap.paragraph_format.space_after = Pt(4)
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        label = f"Table {self.chapter_no}.{self.table_no}"
        _font(cap.add_run(f"{label} "), 10, bold=True)
        _font(cap.add_run(caption), 10)
        self._mark(cap, "T", 1, f"{self.chapter_no}.{self.table_no}", caption)

        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        align = align or ["left"] * len(headers)
        alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                     "right": WD_ALIGN_PARAGRAPH.RIGHT}

        def fill(cell, text, bold=False, col=0):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else alignment[align[col]]
            self._add_inline(p, str(text), size=font_size) if not bold else \
                _font(p.add_run(str(text)), font_size, bold=True)

        for i, h in enumerate(headers):
            fill(table.rows[0].cells[i], h, bold=True, col=i)
        # Repeat the header row when a table spans pages.
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)

        for row in rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                fill(cells[i], value, col=i)
        if widths:
            for row in table.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        after = self.doc.add_paragraph()
        after.paragraph_format.space_after = Pt(4)
        after.paragraph_format.line_spacing = Pt(6)

    def equation(self, omml: str) -> str:
        """Centred Word equation with its number '(C.N)' at the right margin."""
        self.eq_no += 1
        number = f"{self.chapter_no}.{self.eq_no}"
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN / 2), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        _font(p.add_run("\t"), 12)
        p._p.append(parse_xml(f"<m:oMath {M_NS}>{omml}</m:oMath>"))
        _font(p.add_run(f"\t({number})"), 12)
        return f"Eqn. {number}"

    def code(self, text: str, size: float = 8.5) -> None:
        """Monospaced code block, single line spacing, no shading."""
        for i, line in enumerate(text.rstrip("\n").split("\n")):
            p = self.doc.add_paragraph()
            fmt = p.paragraph_format
            fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
            fmt.space_after = Pt(0)
            fmt.space_before = Pt(4) if i == 0 else Pt(0)
            fmt.left_indent = Inches(0.1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _font(p.add_run(line if line else " "), size, name=MONO)
        self.spacer(8)

    # ------------------------------------------------------------ lists of contents
    def toc(self, title: str, kind: str, header: tuple[str, str, str] | None = None) -> None:
        """Heading plus a placeholder that finalize_lists() replaces with the entries."""
        self.front_heading(title, {"C": "Table of Contents", "F": "List of Figures", "T": "List of Tables"}[kind])
        if header:
            p = self.doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(Inches(2.6), WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
            _font(p.add_run(f"{header[0]}\t{header[1]}\t{header[2]}"), 12, bold=True)
        elif kind == "C":
            p = self.doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
            _font(p.add_run("\tPage No."), 12, bold=True)
        self._placeholders[kind] = self.doc.add_paragraph()

    def save(self, path: Path) -> None:
        self.finalize_lists()
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
